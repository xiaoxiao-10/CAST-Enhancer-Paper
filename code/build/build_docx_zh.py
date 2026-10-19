#!/usr/bin/env python
"""Build the Chinese Word manuscript from manuscript_v3_zh.md."""
import os
import re

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
MD = os.path.join(ROOT, "manuscript", "zh", "manuscript_v3_zh.md")
MD_EN = os.path.join(ROOT, "manuscript", "en", "manuscript_v3.md")
FIG = os.path.join(ROOT, "figures", "main")
TAB = os.path.join(ROOT, "tables")
LHM = os.path.join(ROOT, "results", "locus_holdout")
OUT = os.path.join(ROOT, "manuscript", "zh", "CAST-Enhancer_v3_中文稿.docx")

CN = "宋体"
CN_BOLD = "黑体"
LATIN = "Times New Roman"


def set_run_font(run, size=None, bold=None, italic=None, cn=CN, latin=LATIN,
                 color=None):
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:ascii"), latin)
    rf.set(qn("w:hAnsi"), latin)
    rf.set(qn("w:eastAsia"), cn)
    if color is not None:
        run.font.color.rgb = color


def add_rich(p, text, size, bold=False, color=None):
    """Parse **bold** and *italic* segments."""
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        b, i, t = bold, False, part
        if part.startswith("**") and part.endswith("**"):
            b, t = True, part[2:-2]
        elif part.startswith("*") and part.endswith("*"):
            i, t = True, part[1:-1]
        elif part.startswith("`") and part.endswith("`"):
            t = part[1:-1]
            r = p.add_run(t)
            set_run_font(r, size=size - 0.5, latin="Consolas", cn=CN, color=color)
            continue
        r = p.add_run(t)
        set_run_font(r, size=size, bold=b, italic=i, color=color)


def style_para(p, before=0, after=4, line=1.45, indent=None, align=None,
               keep_next=False):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if indent is not None:
        pf.first_line_indent = Cm(indent)
    if align is not None:
        p.alignment = align
    pf.keep_with_next = keep_next


def heading(doc, text, level):
    p = doc.add_paragraph()
    style_para(p, before=(16 if level == 1 else 10),
               after=(6 if level == 1 else 4), line=1.2, keep_next=True)
    r = p.add_run(text)
    set_run_font(r, size=(14 if level == 1 else 12), bold=True, cn=CN_BOLD)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def body(doc, text, size=11):
    p = doc.add_paragraph()
    align = WD_ALIGN_PARAGRAPH.LEFT if "`" in text else WD_ALIGN_PARAGRAPH.JUSTIFY
    style_para(p, after=5, line=1.45, indent=0.78, align=align)
    add_rich(p, text, size)
    return p


def caption(doc, text, size=9, keep_next=False):
    p = doc.add_paragraph()
    style_para(p, before=2, after=10, line=1.2, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
               keep_next=keep_next)
    add_rich(p, text, size)
    return p


def bullet(doc, text, size=11):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.78)
    pf.first_line_indent = Cm(-0.39)
    pf.space_before = Pt(0)
    pf.space_after = Pt(3)
    pf.line_spacing = 1.35
    r = p.add_run("• ")
    set_run_font(r, size=size)
    add_rich(p, text, size)
    return p


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for el, attr in (("w:fldChar", {"w:fldCharType": "begin"}),
                     ("w:instrText", None),
                     ("w:fldChar", {"w:fldCharType": "end"})):
        e = OxmlElement(el)
        if attr:
            for k, v in attr.items():
                e.set(qn(k), v)
        if el == "w:instrText":
            e.set(qn("xml:space"), "preserve")
            e.text = " PAGE "
        run._r.append(e)
    set_run_font(run, size=9)


def set_cell_edge(cell, edge, sz):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    el = OxmlElement("w:" + edge)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), "000000")
    borders.append(el)


def three_line(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge, sz in (("top", 12), ("bottom", 12)):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    for edge in ("left", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def keep_table_together(table):
    rows = table.rows
    for ri, row in enumerate(rows):
        trPr = row._tr.get_or_add_trPr()
        el = OxmlElement("w:cantSplit")
        trPr.append(el)
        if ri < len(rows) - 1:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True


def cell_text(cell, text, size=8.5, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.vertical_alignment = 1
    p = cell.paragraphs[0]
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(2)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.0
    r = p.add_run(str(text))
    set_run_font(r, size=size, bold=bold)


CELL_CN = {
    "mature adipocyte": "成熟脂肪细胞",
    "Endothelial Cells - Aortic": "主动脉内皮细胞",
    "Bronchial Epithelial Cell": "支气管上皮细胞",
    "Olfactory epithelial cells": "嗅上皮细胞",
    "Fibroblast - Cardiac": "心脏成纤维细胞",
    "Fibroblast - Gingival": "牙龈成纤维细胞",
    "CD19+ B Cells (pluriselect)": "CD19+ B 细胞",
    "CD8+ T Cells (pluriselect)": "CD8+ T 细胞",
    "Skeletal Muscle Cells": "骨骼肌细胞",
    "Smooth Muscle Cells - Aortic": "主动脉平滑肌细胞",
    "Mesenchymal Stem Cells - bone marrow": "骨髓间充质干细胞",
}


def add_table1(doc):
    df = pd.read_csv(os.path.join(TAB, "table1_cohort.tsv"), sep="\t")
    caption(doc, "**表 1 | 队列与标签计数。** 20% 位点留出后训练细胞的计数，以及 11 个留出原代细胞类型的计数（文库数、活性阈值、阳性数、阴性数和阳性率）。", size=9, keep_next=True)
    cols = ["细胞类型", "数据划分", "文库数", "活性阈值", "阳性数", "阴性数", "阳性率"]
    widths = [4.4, 2.0, 1.8, 2.0, 1.9, 2.0, 2.1]
    t = doc.add_table(rows=1, cols=len(cols))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    three_line(t)
    hdr = t.rows[0]
    trPr = hdr._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)
    for i, c in enumerate(cols):
        cell = hdr.cells[i]
        cell.width = Cm(widths[i])
        cell_text(cell, c, bold=True)
        set_cell_edge(cell, "bottom", 8)

    def group(text):
        row = t.add_row()
        cell = row.cells[0].merge(row.cells[-1])
        cell.vertical_alignment = 1
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        set_run_font(r, size=8.5, bold=True)

    def data_row(vals):
        cells = t.add_row().cells
        for i, (v, w) in enumerate(zip(vals, widths)):
            cells[i].width = Cm(w)
            cell_text(cells[i], v, align=(WD_ALIGN_PARAGRAPH.LEFT if i < 2 else WD_ALIGN_PARAGRAPH.CENTER))

    group("训练细胞（20% 位点留出后）")
    for row in df[df.block == "development"].itertuples():
        data_row([row.cell, {"train": "训练", "validation": "验证", "test": "测试"}[row.split],
                  row.n_libraries, row.threshold, row.positives, row.negatives,
                  "%.2f%%" % (100 * row.prevalence)])
    group("留出原代细胞类型（标签仅读取一次）")
    for row in df[df.block == "held-out"].itertuples():
        data_row([CELL_CN.get(row.cell, row.cell), "留出", row.n_libraries, row.threshold,
                  row.positives, row.negatives, "%.2f%%" % (100 * row.prevalence)])
    for cell in t.rows[-1].cells:
        set_cell_edge(cell, "bottom", 12)
    keep_table_together(t)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table2(doc):
    perf = pd.read_csv(os.path.join(TAB, "table2_four_strata.tsv"), sep="\t")
    macro = pd.read_csv(os.path.join(LHM, "LH_macro_point.tsv"), sep="\t")
    caption(doc, "**表 2 | 四分层性能。** 三个模型在各分层中的平均 AUPRC 和提升倍数（括号内为逐细胞 lift 均值）；n 和阳性数为逐细胞均值；未见不活跃行使用 9 个至少含 10 个阳性的细胞类型；Δ 为 NT-LoRA 减组成模型的点估计及 95% CI。", size=9, keep_next=True)
    cols = ["位点集合", "训练细胞活性", "n / 阳性数", "组成模型", "DNA-CNN", "NT-LoRA", "ΔAUPRC（95% CI）"]
    widths = [2.4, 2.4, 2.3, 2.1, 2.1, 2.1, 2.8]
    t = doc.add_table(rows=1, cols=len(cols))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    three_line(t)
    hdr = t.rows[0]
    trPr = hdr._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)
    for i, c in enumerate(cols):
        cell = hdr.cells[i]
        cell.width = Cm(widths[i])
        cell_text(cell, c, bold=True)
        set_cell_edge(cell, "bottom", 8)
    order = [("unseen", "all", "未见", "全部"),
             ("unseen", "dev_active", "未见", "活跃"),
             ("seen", "dev_active", "保留", "活跃"),
             ("unseen", "dev_inactive", "未见", "不活跃")]
    for sub, dact, s_cn, d_cn in order:
        base = perf[(perf.subset == sub) & (perf.dact == dact)].iloc[0]
        md = macro[(macro.subset == sub) & (macro.dact == dact)].iloc[0]
        vals = [s_cn, d_cn,
                "%s / %d" % (format(int(round(base.mean_n)), ","), int(round(base.mean_pos)))]
        for m in ("comp", "cnn", "nt"):
            vals.append("%.3f（%.2f×）" % (base[m + "_auprc"], base[m + "_lift"]))
        vals.append("%+.4f（%+.4f, %+.4f）" % (md.point_delta, md.ci_lo, md.ci_hi))
        cells = t.add_row().cells
        for i, (v, w) in enumerate(zip(vals, widths)):
            cells[i].width = Cm(w)
            cell_text(cells[i], v,
                      bold=((sub, dact) == ("unseen", "dev_active")),
                      align=(WD_ALIGN_PARAGRAPH.LEFT if i < 2 else WD_ALIGN_PARAGRAPH.CENTER))
    for cell in t.rows[-1].cells:
        set_cell_edge(cell, "bottom", 12)

    keep_table_together(t)


def add_table3(doc):
    comp = pd.read_csv(os.path.join(LHM, "LH_model_compare_macro.tsv"), sep="\t")
    strong = pd.read_csv(os.path.join(LHM, "LH_nt_vs_strong_macro.tsv"), sep="\t")
    boot = pd.read_csv(os.path.join(LHM, "LH_nt_vs_strong_boot.tsv"), sep="\t")
    caption(doc, "**表 3 | 更强的组成与 motif 基线。** 逻辑回归组成、组成-XGBoost、motif-XGBoost、组成+motif-XGBoost 和 NT-LoRA 在四个分层中的 AUPRC；Δ 为 NT 减组成-XGBoost 的宏观差值及 95% CI。", size=9, keep_next=True)
    cols = ["位点集合", "训练细胞活性", "逻辑回归组成", "组成-XGBoost",
            "Motif-XGBoost", "组成+motif-XGBoost", "NT-LoRA"]
    widths = [2.3, 2.3, 2.4, 2.4, 2.3, 2.6, 1.9]
    t = doc.add_table(rows=1, cols=len(cols))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    three_line(t)
    hdr = t.rows[0]
    trPr = hdr._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)
    for i, c in enumerate(cols):
        cell = hdr.cells[i]
        cell.width = Cm(widths[i])
        cell_text(cell, c, bold=True, size=8.0)
        set_cell_edge(cell, "bottom", 8)
    order = [("unseen", "all", "未见", "全部"),
             ("unseen", "dev_active", "未见", "活跃"),
             ("seen", "dev_active", "保留", "活跃"),
             ("unseen", "dev_inactive", "未见", "不活跃")]
    for sub, dact, s_cn, d_cn in order:
        vals = [s_cn, d_cn]
        for model in ("comp_logit", "comp_xgb", "motif_xgb", "comb_xgb"):
            v = comp[(comp.subset == sub) & (comp.dact == dact) &
                     (comp.model == model)].auprc.iloc[0]
            vals.append("%.3f" % v)
        srow = strong[(strong.subset == sub) & (strong.dact == dact)].iloc[0]
        vals.append("%.3f" % srow.nt)
        cells = t.add_row().cells
        for i, (v, w) in enumerate(zip(vals, widths)):
            cells[i].width = Cm(w)
            cell_text(cells[i], v, size=7.8,
                      bold=((sub, dact) == ("unseen", "dev_active")),
                      align=(WD_ALIGN_PARAGRAPH.LEFT if i < 2 else WD_ALIGN_PARAGRAPH.CENTER))
    for cell in t.rows[-1].cells:
        set_cell_edge(cell, "bottom", 12)
    parts = []
    for sub, dact, s_cn, d_cn in order:
        srow = strong[(strong.subset == sub) & (strong.dact == dact)].iloc[0]
        brow = boot[(boot.subset == sub) & (boot.dact == dact)].iloc[0]
        parts.append("%s %s：%+.4f（%+.4f 至 %+.4f）" %
                     (s_cn, d_cn, srow.nt_minus_comp_xgb,
                      brow.comp_xgb_ci_lo, brow.comp_xgb_ci_hi))
    row = t.add_row()
    cell = row.cells[0].merge(row.cells[-1])
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    add_rich(p, "Δ（NT − 组成-XGBoost）：" + "；".join(parts), size=8.0)
    set_cell_edge(cell, "bottom", 12)
    keep_table_together(t)


def add_references(doc):
    text = open(MD_EN, encoding="utf-8").read()
    sec = text.split("## References", 1)[1].split("\n---", 1)[0]
    refs = [ln.strip() for ln in sec.splitlines() if re.match(r"^\d+\.\s", ln.strip())]
    for ref in refs:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.65)
        pf.first_line_indent = Cm(-0.65)
        pf.space_before = Pt(0)
        pf.space_after = Pt(2)
        pf.line_spacing = 1.12
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_rich(p, ref, 9)


def smart_join(lines):
    out = ""
    for ln in lines:
        if not out:
            out = ln
        elif re.search(r"[A-Za-z0-9]$", out) and re.match(r"^[A-Za-z0-9]", ln):
            out += " " + ln
        else:
            out += ln
    return out


def main():
    lines = open(MD, encoding="utf-8").read().splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.4)
    sec.top_margin = Cm(2.3)
    sec.bottom_margin = Cm(2.3)
    add_page_number(sec)

    normal = doc.styles["Normal"]
    normal.font.name = LATIN
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CN)

    i = 0
    in_legends = False
    legends_heading_done = False
    expect_caption = None
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph()
            style_para(p, before=6, after=4, line=1.25, align=WD_ALIGN_PARAGRAPH.CENTER)
            r = p.add_run(line[2:].strip())
            set_run_font(r, size=17, bold=True, cn=CN_BOLD)
            r.font.color.rgb = RGBColor(0, 0, 0)
            p2 = doc.add_paragraph()
            style_para(p2, before=0, after=2, line=1.1, align=WD_ALIGN_PARAGRAPH.CENTER)
            r2 = p2.add_run("中文稿 | 2026 年 10 月 19 日")
            set_run_font(r2, size=9.5, color=RGBColor(80, 80, 80))
            p3 = doc.add_paragraph()
            style_para(p3, before=0, after=12, line=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
            r3 = p3.add_run("说明：正文、图注与表头为中文；图中标签与参考文献保留英文，以便与原始图件及国际期刊版本对照。")
            set_run_font(r3, size=9, color=RGBColor(80, 80, 80))
            i += 1
            continue
        if line.startswith("### "):
            heading(doc, line[4:].strip(), 2)
            i += 1
            continue
        if line.startswith("## "):
            title = line[3:].strip()
            if title == "图注":
                in_legends = True
                i += 1
                continue
            if title == "参考文献":
                heading(doc, title, 1)
                i += 1
                continue
            heading(doc, title, 1)
            i += 1
            continue
        if line.startswith("![Figure"):
            m = re.search(r"\((.*?)\)", line)
            num = re.search(r"Figure\s+(\d+)", line).group(1)
            path = os.path.join(FIG, os.path.basename(m.group(1)))
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(path, width=Inches(6.35))
            alt = {
                "1": "图 1：位点留出与细胞类型留出的实验设计。",
                "2": "图 2：五个组成、motif 与 NT 模型在四个位点分层中的宏观 AUPRC。",
                "3": "图 3：NT 与组成-XGBoost 的逐细胞散点和宏观森林图。",
                "4": "图 4：预训练与随机初始化 NT 的对照。",
            }.get(num, "图件")
            docPr = doc.inline_shapes[-1]._inline.docPr
            docPr.set("descr", alt)
            docPr.set("title", alt)
            expect_caption = True
            i += 1
            continue
        if line.strip() == "<!--TABLES-->":
            hp = heading(doc, "表格", 1)
            hp.paragraph_format.page_break_before = False
            add_table1(doc)
            add_table2(doc)
            add_table3(doc)
            i += 1
            continue
        if line.strip() == "<!--REFERENCES-->":
            add_references(doc)
            i += 1
            continue
        if in_legends:
            # Only supplementary figure legends are repeated here; main figure and table
            # captions already sit next to their objects.
            if line.startswith("**补充图"):
                if not legends_heading_done:
                    heading(doc, "补充图注", 1)
                    legends_heading_done = True
                caption(doc, line.strip(), size=9)
            i += 1
            continue
        if line.startswith("**图 ") or line.startswith("**表 "):
            caption(doc, line.strip(), size=9)
            expect_caption = False
            i += 1
            continue
        if line.startswith("- "):
            bullet(doc, line[2:].strip())
            i += 1
            continue
        # collect a paragraph
        buf = [line.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(("#", "![", "- ", "<!--")):
            buf.append(lines[i].strip())
            i += 1
        text = smart_join(buf)
        if text.startswith("**关键词"):
            p = doc.add_paragraph()
            style_para(p, before=2, after=8, line=1.35, align=WD_ALIGN_PARAGRAPH.LEFT)
            add_rich(p, text, 10.5)
        else:
            body(doc, text)

    doc.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
